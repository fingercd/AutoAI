"""Batch prediction workbook (.xlsx) and per-class metric projection coverage."""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from backend.app.runs.artifacts import ARTIFACT_CATALOG, RunArtifactWriter
from backend.app.runs.batch_projection import project_model_comparison
from backend.app.runs.contracts import Principal
from backend.app.runs.prediction_export import (
    MODEL_DISPLAY_NAMES,
    PredictionExportError,
    build_prediction_workbook,
)
from backend.app.runs.repository import RunRepository

HEADER = ["Index", "Label", "Sample_ID", "划分"]


def _repository(tmp_path: Path) -> RunRepository:
    repository = RunRepository(tmp_path / "runs.sqlite3")
    repository.initialize()
    return repository


def _batch(repository: RunRepository, *, principal: Principal = Principal()):
    return repository.create_batch_queued(
        dataset_id="dataset-1",
        test_dataset_id=None,
        legacy_data_path=None,
        config={"normalization": "zscore", "evaluation_strategy": "stratified_holdout", "split_mode": "stratified_holdout"},
        model_types=["pls_da", "svm"],
        repeat_count=1,
        base_seed=73,
        dataset_snapshot={"name": "fixture.csv", "sha256": "a" * 64},
        principal=principal,
    )


# (index, Sample_ID, true_label, pred_label, "prob_A,prob_B") for every record.
_PRIMARY = [
    ("1", "S1", "A", "A", "0.90,0.10"),
    ("2", "S2", "B", "B", "0.20,0.80"),
    ("3", "S3", "A", "B", "0.45,0.55"),
    ("4", "S4", "B", "B", "0.10,0.90"),
]
_EXTERNAL = [
    ("5", "E1", "A", "A", "0.70,0.30"),
    ("6", "E2", "B", "A", "0.51,0.49"),
]


def _record_rows(*, flip_sample_three: bool = False) -> list[tuple[str, str, str, str, str, str, str]]:
    """Return (dataset, split, index, Sample_ID, true, pred, probs) rows."""
    rows: list[tuple[str, str, str, str, str, str, str]] = []
    for dataset, split, source in (("primary", "train", _PRIMARY[:2]), ("primary", "valid", _PRIMARY[2:3]), ("primary", "test", _PRIMARY[3:]), ("external_test", "external_test", _EXTERNAL)):
        for index, sample_id, true_label, predicted, probabilities in source:
            if flip_sample_three and index == "3":
                predicted, probabilities = "A", "0.62,0.38"
            rows.append((dataset, split, index, sample_id, true_label, predicted, probabilities))
    return rows


def _write_run(
    run_dir: Path,
    run_id: str,
    *,
    precision: float,
    recall: float,
    include_all_predictions: bool = True,
    flip_sample_three: bool = False,
) -> None:
    writer = RunArtifactWriter(run_dir)
    writer.write_json("metrics.json", {"test": {
        "accuracy": 0.75, "balanced_accuracy": 0.75, "macro_f1": 0.5, "weighted_f1": 0.5,
        "classification_report": {
            "A": {"precision": precision, "recall": recall, "f1": 0.5, "support": 2},
            "B": {"precision": 1 - precision, "recall": 1 - recall, "f1": 0.25, "support": 2},
            "macro avg": {"precision": 0.5, "recall": 0.5, "f1": 0.4, "support": 4},
            "weighted avg": {"precision": 0.5, "recall": 0.5, "f1": 0.4, "support": 4},
        },
        "confusion_matrix": [[3, 1], [1, 3]],
    }})
    writer.write_json("split.json", [{
        "fold_index": 1, "train_sample_ids": ["S1", "S2"],
        "valid_sample_ids": ["S3"], "test_sample_ids": ["S4"],
    }])
    writer.write_json("label_map.json", {"0": "A", "1": "B"})
    # 历史 predictions.csv：有 dataset 列，但没有 index/split 列。
    writer.write_bytes("predictions.csv", (
        "dataset,fold_index,Sample_ID,true_label,pred_label,prob_A,prob_B\n"
        "test,1,S4,B,B,0.10,0.90\n"
        "external_test,1,E1,A,A,0.70,0.30\n"
        "external_test,1,E2,B,A,0.51,0.49\n"
    ).encode("utf-8"))
    if include_all_predictions:
        lines = ["dataset,split,fold_index,index,Sample_ID,true_label,pred_label,prob_A,prob_B"]
        lines += [
            f"{dataset},{split},1,{index},{sample_id},{true_label},{predicted},{probabilities}"
            for dataset, split, index, sample_id, true_label, predicted, probabilities in _record_rows(flip_sample_three=flip_sample_three)
        ]
        writer.write_bytes("all_predictions.csv", ("\n".join(lines) + "\n").encode("utf-8"))
    writer.finalize(run_id=run_id)


def _finished_batch(
    tmp_path: Path,
    *,
    precision: float,
    recall: float,
    include_all_predictions: bool = True,
    flip_model: str | None = None,
):
    repository = _repository(tmp_path)
    batch, _records = _batch(repository)
    now = datetime.now(timezone.utc)
    while True:
        claimed = repository.claim_next(worker_id="test", now=now)
        if claimed is None:
            break
        _write_run(
            tmp_path / "runs" / claimed.run_id,
            claimed.run_id,
            precision=precision,
            recall=recall,
            include_all_predictions=include_all_predictions,
            flip_sample_three=str(claimed.config.get("model_type")) == flip_model,
        )
        repository.finish_success(claimed.run_id, claim_token=claimed.claim_token or "", now=now)
    records = repository.list_batch_runs_scoped(batch.batch_id, principal=Principal())
    comparison = project_model_comparison(
        batch_id=batch.batch_id, records=records, run_dir_for=lambda run_id: tmp_path / "runs" / run_id,
    )
    assert comparison["comparable"] is True, comparison.get("reason")
    return comparison, records


def _workbook(tmp_path: Path, comparison, records):
    content = build_prediction_workbook(
        batch_id=comparison["batch_id"], records=records,
        run_dir_for=lambda run_id: tmp_path / "runs" / run_id, comparison=comparison,
    )
    return load_workbook(BytesIO(content))


def test_all_predictions_artifact_is_public_but_optional() -> None:
    entry = ARTIFACT_CATALOG["all_predictions.csv"]
    assert entry["downloadable"] is True
    # required=False keeps historical Manifests valid without the new file.
    assert entry["required"] is False


def test_workbook_keeps_identity_columns_and_one_column_per_model(tmp_path: Path) -> None:
    comparison, records = _finished_batch(tmp_path, precision=0.9, recall=0.8)
    workbook = _workbook(tmp_path, comparison, records)
    names = workbook.sheetnames
    assert names[0] == "预测类别"
    assert names[1].startswith("预测概率")

    predicted = workbook[names[0]]
    assert [cell.value for cell in predicted[1]] == [*HEADER, "PLS-DA", "SVM"]
    rows = list(predicted.iter_rows(min_row=2, values_only=True))
    assert [row[0] for row in rows] == ["1", "2", "3", "4", "5", "6"]
    assert [row[1] for row in rows] == ["A", "B", "A", "B", "A", "B"]
    assert [row[2] for row in rows] == ["S1", "S2", "S3", "S4", "E1", "E2"]
    assert [row[3] for row in rows] == ["train", "train", "valid", "test", "external_test", "external_test"]
    assert [row[4] for row in rows] == ["A", "B", "B", "B", "A", "A"]
    assert [row[5] for row in rows] == ["A", "B", "B", "B", "A", "A"]

    probabilities = workbook[names[1]]
    assert [cell.value for cell in probabilities[1]] == [*HEADER, "PLS-DA", "SVM"]
    assert [row[4] for row in probabilities.iter_rows(min_row=2, values_only=True)] == [
        "0.900000,0.100000", "0.200000,0.800000", "0.450000,0.550000", "0.100000,0.900000",
        "0.700000,0.300000", "0.510000,0.490000",
    ]
    for row in probabilities.iter_rows(min_row=2):
        for cell in row[4:]:
            assert cell.number_format == "@"
    # 概率列头注释说明类别顺序，避免事后猜测每个数字对应哪个类别。
    assert probabilities.cell(row=1, column=5).comment is not None


def test_workbook_reports_each_model_own_predictions(tmp_path: Path) -> None:
    comparison, records = _finished_batch(tmp_path, precision=0.9, recall=0.8, flip_model="svm")
    rows = list(_workbook(tmp_path, comparison, records)["预测类别"].iter_rows(min_row=2, values_only=True))
    # Sample S3（index 3）是两个模型唯一给出不同预测的记录。
    assert rows[2][4] == "B" and rows[2][5] == "A"


def test_workbook_falls_back_to_test_rows_for_historical_runs(tmp_path: Path) -> None:
    comparison, records = _finished_batch(tmp_path, precision=0.9, recall=0.8, include_all_predictions=False)
    rows = list(_workbook(tmp_path, comparison, records)["预测类别"].iter_rows(min_row=2, values_only=True))
    # 历史 Run 只有 test 预测产物：只导出这些行，不补造 train/valid；历史文件没有
    # index 列，Index 留空而不是编造序号，行序按“主数据在前、独立测试集在后”稳定排序。
    assert [row[2] for row in rows] == ["S4", "E1", "E2"]
    assert [row[3] for row in rows] == ["test", "external_test", "external_test"]
    assert all(row[0] is None for row in rows)


def test_workbook_rejects_uncomparable_batches(tmp_path: Path) -> None:
    with pytest.raises(PredictionExportError):
        build_prediction_workbook(
            batch_id="batch_x",
            records=[],
            run_dir_for=lambda run_id: tmp_path,
            comparison={"comparable": False, "reason": "批次中尚无成功且结果完整的模型"},
        )


def test_class_metrics_separate_precision_from_recall(tmp_path: Path) -> None:
    comparison, _records = _finished_batch(tmp_path, precision=0.9, recall=0.55)
    block = comparison["class_metrics"]
    assert block["status"] == "ready"
    assert block["labels"] == ["A", "B"]
    for row in block["rows"]:
        first, second = row["values"]
        assert first["precision"] == pytest.approx(0.9)
        assert first["recall"] == pytest.approx(0.55)
        assert first["f1"] == pytest.approx(0.5)
        assert first["support"] == 2
        assert second["precision"] == pytest.approx(0.1)
        assert second["recall"] == pytest.approx(0.45)


def test_model_display_names_match_figure_names() -> None:
    from backend.app.runs.comparison_figures import NAMES

    assert MODEL_DISPLAY_NAMES == NAMES


def test_predictions_xlsx_route_is_registered() -> None:
    from backend.app.main import app

    paths = {getattr(route, "path", "") for route in app.routes}
    assert "/api/training/batches/{batch_id}/predictions.xlsx" in paths
