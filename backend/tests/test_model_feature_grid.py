"""Cartesian product batches, scoped comparison groups, CV and non-colliding exports."""
from dataclasses import replace
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app.contracts import TrainingBatchRequest, TrainingConfigValidationError
from backend.app.routers.batches import _batch_config
from backend.app.runs.repository import RunRepository
from backend.app.runs.contracts import Principal
from backend.app.runs.batch_projection import batch_status_payload, project_model_comparison
from backend.app.runs.comparison_figures import figure_data, figure_specs
from backend.app.runs.prediction_export import build_prediction_workbook
from backend.tests.test_training_batches import _write_comparison_artifacts


def grid(repository, models=None, features=None, mode='stratified_holdout'):
    request=TrainingBatchRequest(model_types=models or ['pls_da','svm'],feature_schemes=features or ['bin_3','pca_components_5'],
        config={'experiment_version':'word-0904','split_mode':mode})
    types,config,_=_batch_config(request,has_external_test='external' in mode)
    config['expected_outer_folds']=24 if 'leave_one' in mode else 1
    return repository.create_batch_queued(dataset_id='d',test_dataset_id=None,legacy_data_path=None,
        config=config,model_types=types,repeat_count=1,base_seed=42,dataset_snapshot={'sha256':'a'*64,'name':'fixture.csv'})


def repo(tmp_path):
    result=RunRepository(tmp_path/'runs.sqlite3');result.initialize();return result


def test_cartesian_grid_and_cv_workload_deduplicate_and_keep_native_model_identity(tmp_path):
    repository=repo(tmp_path)
    batch,runs=grid(repository,models=['pls_da','svm','svm'],features=['bin_3','pca_components_5','bin_3'],mode='leave_one_sample_id_cv')
    assert len(runs)==4
    assert {(r.config['model_type'],r.config['feature_scheme']) for r in runs}=={(m,s) for m in ['pls_da','svm'] for s in ['bin_3','pca_components_5']}
    assert len({r.batch_model_order for r in runs})==4
    assert {r.split_seed for r in runs}=={42}
    status=batch_status_payload(batch.batch_id,runs)
    assert status['workload']['total_outer_folds']==96
    assert status['workload']['model_count']==2 and status['workload']['feature_count']==2
    assert all(row['expected_outer_folds']==24 for row in status['runs'])
    pending=project_model_comparison(batch_id=batch.batch_id,records=runs,run_dir_for=lambda rid:tmp_path/rid)
    assert not pending['comparable']
    assert pending['evaluation']['primary_aggregation']=='pooled_oof'
    assert pending['combination_status']['workload']['total_outer_folds']==96
    repository.cancel_batch_scoped(batch.batch_id,now=datetime.now(timezone.utc),principal=Principal())
    assert all(r.state=='cancelled' for r in repository.list_batch_runs_scoped(batch.batch_id,principal=Principal()))


def test_incompatible_pairs_are_explicit_and_do_not_replace_feature_scheme(tmp_path):
    repository=repo(tmp_path)
    batch,runs=grid(repository,models=['pls_da','cnn1d'])
    assert len(runs)==3
    assert batch.config['inapplicable_combinations']==[{'model_type':'cnn1d','feature_scheme':'pca_components_5','reason':'CNN 不支持 PCA 输入'}]
    assert {(r.config['model_type'],r.config['feature_scheme']) for r in runs}=={('pls_da','bin_3'),('pls_da','pca_components_5'),('cnn1d','bin_3')}
    with pytest.raises(TrainingConfigValidationError,match='没有可运行'):
        _batch_config(TrainingBatchRequest(model_types=['cnn1d'],feature_schemes=['pca_components_5'],config={'experiment_version':'word-0904'}),has_external_test=False)


def test_legacy_batches_do_not_invent_unrecorded_cv_workload(tmp_path):
    repository=repo(tmp_path)
    batch,runs=repository.create_batch_queued(dataset_id='d',test_dataset_id=None,legacy_data_path=None,
        config={'split_mode':'leave_one_sample_id_cv'},model_types=['pls_da'],repeat_count=2,base_seed=42,dataset_snapshot={})
    assert 'workload' not in batch_status_payload(batch.batch_id,runs)


@pytest.mark.parametrize('config,features',[({'experiment_version':'word-0904','feature_scheme':'full'},['bin_3']),({'experiment_version':'word-0904','training_profile':'full'},['bin_3']),({},['bin_3']),({'experiment_version':'word-0904'},['bad'])])
def test_invalid_grid_config_fails_before_insertion(config,features):
    with pytest.raises(TrainingConfigValidationError):
        _batch_config(TrainingBatchRequest(model_types=['svm'],feature_schemes=features,config=config),has_external_test=False)


def test_grid_run_limit_is_product_not_just_model_count():
    from backend.app.feature_engineering import ALL_SCHEMES, MODELS
    with pytest.raises(TrainingConfigValidationError,match='不能超过 100'):
        _batch_config(TrainingBatchRequest(model_types=sorted(MODELS),feature_schemes=[key for key,_ in ALL_SCHEMES],config={'experiment_version':'word-0904'}),has_external_test=False)


def finished_grid(tmp_path):
    repository=repo(tmp_path);batch,runs=grid(repository)
    finished=[]
    for i,record in enumerate(runs):
        _write_comparison_artifacts(tmp_path/record.run_id,record.run_id,score=.5+i*.1,sample_suffix='same')
        finished.append(replace(record,state='succeeded'))
    return batch,finished


def test_projection_never_collapses_features_into_historical_repeats(tmp_path):
    batch,runs=finished_grid(tmp_path)
    data=project_model_comparison(batch_id=batch.batch_id,records=runs,run_dir_for=lambda rid:tmp_path/rid)
    assert data['comparable'] and len(data['models'])==4
    assert 'run_comparisons' not in data
    assert {row['comparison_id'] for row in data['models']}=={f'{m}__{s}' for m in ['pls_da','svm'] for s in ['bin_3','pca_components_5']}
    assert all(len(row['run_ids'])==1 and row['successful_repeats']==1 for row in data['models'])
    matrix=figure_data(data,kind='features')
    assert matrix['columns']==['PLS-DA','SVM'] and len(matrix['rows'])==2
    assert matrix['values']==[[.5,.7],[.6,.8]]
    assert not any(any(line) for line in matrix['selected'])
    assert all(len(g['models'])==2 for g in data['scheme_comparisons'])
    with pytest.raises(ValueError,match='必须指定'):
        figure_data(data,kind='matrix',model='pls_da')
    with pytest.raises(ValueError,match='不包含'):
        figure_data(data,kind='matrix',model='pls_da',feature_scheme='bad')
    detail=figure_data(data,kind='matrix',model='pls_da',feature_scheme='bin_3',matrix_mode='count')
    assert sum(map(sum,detail['values']))==4
    specs=figure_specs(data)
    assert len({s['id'] for s in specs})==len(specs)
    assert any(s.get('feature_scheme')=='bin_3' for s in specs)
    workbook=load_workbook(BytesIO(build_prediction_workbook(batch_id=batch.batch_id,records=runs,run_dir_for=lambda rid:tmp_path/rid,comparison=data)))
    headers=[cell.value for cell in workbook.worksheets[0][1]]
    assert len(headers)==8 and len(set(headers))==8
    assert {'PLS-DA · BIN3','PLS-DA · PCA5','SVM · BIN3','SVM · PCA5'}.issubset(headers)


def test_cross_scheme_comparison_still_blocks_different_snapshot(tmp_path):
    batch,runs=finished_grid(tmp_path)
    runs[-1]=replace(runs[-1],dataset_snapshot={'sha256':'different'})
    data=project_model_comparison(batch_id=batch.batch_id,records=runs,run_dir_for=lambda rid:tmp_path/rid)
    assert not data['comparable'] and '不一致' in data['reason']


def test_all_inapplicable_column_has_reason_and_never_fills_zero(tmp_path):
    repository=repo(tmp_path)
    batch,runs=grid(repository,models=['cnn1d'],features=['pca_components_5','bin_3'])
    assert len(runs)==1
    record=runs[0]
    _write_comparison_artifacts(tmp_path/record.run_id,record.run_id,score=.8,sample_suffix='same')
    data=project_model_comparison(batch_id=batch.batch_id,records=[replace(record,state='succeeded')],run_dir_for=lambda rid:tmp_path/rid)
    matrix=figure_data(data,kind='features')
    assert matrix['values']==[[None],[.8]]
    assert 'CNN' in matrix['reasons'][0][0]
    assert data['scheme_comparisons'][0]['state']=='not_applicable'


@pytest.mark.parametrize('mode',['stratified_holdout','leave_one_sample_id_cv','external_test_holdout','leave_one_sample_id_cv_with_external_test'])
def test_grid_all_four_request_modes_keep_default_ratios(mode):
    req=TrainingBatchRequest(model_types=['pls_da','svm'],feature_schemes=['bin_3','pca_components_5'],config={'experiment_version':'word-0904','split_mode':mode,'split_train':8,'split_valid':1 if mode=='stratified_holdout' else 2,**({'split_test':1} if mode=='stratified_holdout' else {})})
    _,config,_=_batch_config(req,has_external_test='external' in mode)
    assert len(config['combinations'])==4 and config['split_mode']==mode


@pytest.mark.parametrize('mode',['stratified_holdout','leave_one_sample_id_cv','external_test_holdout','leave_one_sample_id_cv_with_external_test'])
def test_training_grid_four_modes_keep_oof_and_external_test_separate(tmp_path,monkeypatch,mode):
    from backend.app import training as t
    from backend.tests.test_six_feature_schemes import grouped_frame
    main=tmp_path/'main.csv';grouped_frame(12).to_csv(main,index=False)
    request=TrainingBatchRequest(model_types=['pls_da','svm'],feature_schemes=['bin_3','pca_components_5'],config={'experiment_version':'word-0904','split_mode':mode})
    models,config,_=_batch_config(request,has_external_test='external' in mode)
    config['expected_outer_folds']=12 if 'leave_one' in mode else 1
    if 'external' in mode:
        ext=tmp_path/'external.csv';grouped_frame(4,100).to_csv(ext,index=False)
        config.update(test_data_path=str(ext),test_dataset_sha256='b'*64)
    repository=repo(tmp_path)
    batch,records=repository.create_batch_queued(dataset_id=None,test_dataset_id=None,legacy_data_path=str(main),config=config,model_types=models,repeat_count=1,base_seed=42,dataset_snapshot={'sha256':'a'*64})
    monkeypatch.setattr(t,'RUNS_DIR',tmp_path/'artifacts')
    done=[]
    for record in records:
        result=t.train_model(main,record.config,run_id=record.run_id)
        assert sum(map(sum,result['metrics']['test']['confusion_matrix']))==(8 if 'external' in mode else 24 if 'leave_one' in mode else 4)
        done.append(replace(record,state='succeeded'))
    data=project_model_comparison(batch_id=batch.batch_id,records=done,run_dir_for=lambda rid:t.RUNS_DIR/rid)
    assert data['comparable'],data['reason']
    assert len(data['models'])==4
    assert data['evaluation']['primary_aggregation']==('pooled_oof' if mode=='leave_one_sample_id_cv' else 'direct_external_test' if mode=='leave_one_sample_id_cv_with_external_test' else 'direct')
    for group in data['scheme_comparisons']:
        for row in group['models']:
            exp=row['experiment']
            assert len(exp['fold_configurations'])==(12 if 'leave_one' in mode else 1)
            assert (exp['selected_configuration'] is None)==(mode=='leave_one_sample_id_cv')
            assert all(fold['scheme_id']==group['feature_scheme'] for fold in exp['fold_configurations'])
    workbook=load_workbook(BytesIO(build_prediction_workbook(batch_id=batch.batch_id,records=done,run_dir_for=lambda rid:t.RUNS_DIR/rid,comparison=data)))
    assert all(sheet.max_column==8 and sheet.max_row==(33 if 'external' in mode else 25) for sheet in workbook.worksheets)
