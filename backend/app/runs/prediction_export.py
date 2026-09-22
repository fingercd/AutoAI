"""Batch prediction workbook (.xlsx) export for the classic comparison page.

The classic comparison page offers an Excel download next to the archive ZIP.
This module turns the Manifest-verified prediction artifacts of every successful
sub-Run into one workbook with two sheets:

* ``预测类别``  – ``Index`` / ``Label`` / ``Sample_ID`` / ``划分`` plus one column
  per model, each cell holding that model's predicted class for the record.
* ``预测概率``  – identical layout, but each model cell holds the per-class
  probabilities in ``label_map`` order, comma separated (they sum to 1).

Which rows are covered depends on the evaluation strategy, and the answer is
always taken from a verified artifact rather than recomputed here:

* ``stratified_holdout`` / ``external_test_holdout`` read ``all_predictions.csv``,
  which the trainer writes by running the final fitted model once over every
  record; the ``split`` column then carries train/valid/test (external rows are
  ``external_test``).
* Cross-validation strategies have no single final model, so the same file holds
  the pooled out-of-fold rows – each record appears exactly once, as the test of
  its own fold, and ``split`` is ``test``.

Historical batches trained before this file existed fall back to
``predictions.csv`` (test/OOF rows only); nothing is fabricated and nothing is
retrained.  Server paths never enter the workbook.
"""

from __future__ import annotations

import csv
import json
import math
from io import BytesIO
from pathlib import Path
from typing import Any, Callable, Iterable

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .artifacts import ManifestCorruptError, RunArtifactWriter
from .contracts import RunRecord

PREDICTION_SHEET = "预测类别"
PROBABILITY_SHEET = "预测概率"
METADATA_COLUMNS = ("Index", "Label", "Sample_ID", "划分")
PREDICTION_FILE = "all_predictions.csv"
FALLBACK_PREDICTION_FILE = "predictions.csv"
PROBABILITY_DECIMALS = 6
# Mirrors comparison_figures.NAMES; the frontend keeps the same short names.
# test_prediction_export asserts the two tables cannot drift apart silently.
MODEL_DISPLAY_NAMES = {
    "pls_da": "PLS-DA",
    "logistic_regression": "Elastic Net",
    "svm": "SVM",
    "random_forest": "Random Forest",
    "xgboost": "XGBoost",
    "cnn1d": "1D-CNN",
    "spls_da": "sPLS-DA",
    "pca_svm": "PCA-SVM",
}
EXCEL_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_INVALID_SHEET_CHARS = str.maketrans({char: "_" for char in ":\\/?*[]"})


class PredictionExportError(ValueError):
    """Raised when no trustworthy prediction rows can be exported."""


def _verified_csv(run_dir: Path, run_id: str, name: str) -> list[dict[str, str]] | None:
    """Read one prediction CSV only when the Manifest verifies its integrity."""
    writer = RunArtifactWriter(run_dir)
    try:
        _manifest, descriptors = writer.descriptors(run_id=run_id)
    except (FileNotFoundError, ManifestCorruptError):
        return None
    descriptor = next((item for item in descriptors if item.get("name") == name), None)
    if not isinstance(descriptor, dict) or descriptor.get("integrity") != "ok":
        return None
    try:
        with (run_dir / name).open("r", encoding="utf-8-sig", newline="") as handle:
            return [row for row in csv.DictReader(handle) if any((value or "").strip() for value in row.values())]
    except (OSError, UnicodeDecodeError, csv.Error):
        return None


def _verified_json(run_dir: Path, run_id: str, name: str) -> Any | None:
    writer = RunArtifactWriter(run_dir)
    try:
        _manifest, descriptors = writer.descriptors(run_id=run_id)
    except (FileNotFoundError, ManifestCorruptError):
        return None
    descriptor = next((item for item in descriptors if item.get("name") == name), None)
    if not isinstance(descriptor, dict) or descriptor.get("integrity") != "ok":
        return None
    try:
        return json.loads((run_dir / name).read_text(encoding="utf-8-sig"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError):
        return None


def _numeric(value: str) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return math.inf


def _column_entries(rows: list[dict[str, str]]) -> list[tuple[str, dict[str, str]]]:
    """Attach a stable key to every prediction row, preserving file order.

    ``all_predictions.csv`` carries ``dataset`` + ``index``, so records are keyed
    by those two.  Historical ``predictions.csv`` has no ``index`` column, so it
    falls back to ``Sample_ID`` plus the row position – the predictor writes rows
    in a deterministic fold order, which keeps the columns aligned.
    """
    entries: list[tuple[str, dict[str, str]]] = []
    seen: set[str] = set()
    for row in rows:
        dataset = str(row.get("dataset") or "").strip() or "primary"
        index = str(row.get("index") or "").strip()
        if index:
            key = f"{dataset}|index:{index}"
        else:
            sample_id = str(row.get("Sample_ID") or "").strip()
            key = f"{dataset}|row:{sample_id}#{len(entries)}"
        if key in seen:
            continue
        seen.add(key)
        entries.append((key, row))
    return entries


def _order_key(key: str) -> tuple[int, int, float, str]:
    dataset, _, payload = key.partition("|")
    group = 1 if dataset == "external_test" else 0
    if payload.startswith("index:"):
        return (group, 0, _numeric(payload[len("index:"):]), "")
    return (group, 1, 0.0, payload)


def _labels_for(run_dir: Path, run_id: str) -> list[str]:
    label_map = _verified_json(run_dir, run_id, "label_map.json")
    if not isinstance(label_map, dict):
        return []
    try:
        return [str(label_map[key]) for key in sorted(label_map, key=lambda key: int(key))]
    except (TypeError, ValueError):
        return []


def _sheet_name(base: str, labels: Iterable[str]) -> str:
    suffix = ",".join(str(label) for label in labels)
    name = f"{base}（{suffix}）" if suffix else base
    return name.translate(_INVALID_SHEET_CHARS)[:31]


def _write_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[Any]], *, freeze: bool = True):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2" if freeze else None
    widths = [max(10, min(28, len(str(header)) * 2 + 6)) for header in headers]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    return sheet


def _probability_text(row: dict[str, str], labels: list[str]) -> str:
    parts: list[str] = []
    for label in labels:
        raw = row.get(f"prob_{label}")
        try:
            parts.append(f"{float(raw):.{PROBABILITY_DECIMALS}f}")
        except (TypeError, ValueError):
            parts.append("")
    return ",".join(parts)


def build_prediction_workbook(
    *,
    batch_id: str,
    records: list[RunRecord],
    run_dir_for: Callable[[str], Path],
    comparison: dict[str, Any],
) -> bytes:
    """Build the two-sheet prediction workbook for one comparable batch."""
    if not comparison.get("comparable"):
        raise PredictionExportError(comparison.get("reason") or "当前批次结果不可比较，无法导出预测明细")
    models = list(comparison.get("models") or [])
    if not models:
        raise PredictionExportError("当前批次没有可比较的成功模型")
    by_run_id = {record.run_id: record for record in records}
    columns: list[dict[str, Any]] = []
    used_names: set[str] = set()
    for model in models:
        model_type = str(model.get("model_type") or "")
        run_ids = list(model.get("run_ids") or [])
        record = by_run_id.get(run_ids[0]) if run_ids else None
        if record is None:
            continue
        run_dir = run_dir_for(record.run_id)
        labels = _labels_for(run_dir, record.run_id)
        if not labels:
            continue
        rows = _verified_csv(run_dir, record.run_id, PREDICTION_FILE)
        source = PREDICTION_FILE
        if rows is None:
            # Historical Run without the全量明细 artifact: fall back to the
            # test/OOF prediction rows it does have.
            rows = _verified_csv(run_dir, record.run_id, FALLBACK_PREDICTION_FILE)
            source = FALLBACK_PREDICTION_FILE
        if not rows:
            continue
        by_key = {}
        for key, row in _column_entries(rows):
            by_key.setdefault(key, row)
        display = MODEL_DISPLAY_NAMES.get(model_type, model_type)
        if display in used_names:
            display = f"{display}（{model_type}）"
        used_names.add(display)
        columns.append({"model_type": model_type, "display": display, "labels": labels, "rows": by_key, "source": source})
    if not columns:
        raise PredictionExportError("成功子 Run 缺少通过完整性校验的预测明细或类别映射")

    labels = list(columns[0]["labels"])
    keys = sorted(set().union(*(set(column["rows"]) for column in columns)), key=_order_key)
    if not keys:
        raise PredictionExportError("预测明细中没有可导出的记录")
    metadata: dict[str, tuple[Any, str, str, str]] = {}
    for key in keys:
        row = next((column["rows"].get(key) for column in columns if column["rows"].get(key)), {})
        dataset = key.partition("|")[0]
        split = str(row.get("split") or "").strip() or ("external_test" if dataset == "external_test" else "test")
        metadata[key] = (
            row.get("index", ""),
            str(row.get("true_label") or "").strip(),
            str(row.get("Sample_ID") or "").strip(),
            split,
        )

    workbook = Workbook()
    workbook.remove(workbook.active)
    probability_title = _sheet_name(PROBABILITY_SHEET, labels)
    _write_sheet(
        workbook,
        PREDICTION_SHEET,
        [*METADATA_COLUMNS, *(column["display"] for column in columns)],
        [[*metadata[key], *((column["rows"].get(key) or {}).get("pred_label", "") for column in columns)] for key in keys],
    )
    probability_rows = [
        [*metadata[key], *(_probability_text(column["rows"].get(key) or {}, column["labels"]) for column in columns)]
        for key in keys
    ]
    sheet = _write_sheet(
        workbook,
        probability_title,
        [*METADATA_COLUMNS, *(column["display"] for column in columns)],
        probability_rows,
    )
    for row in sheet.iter_rows(min_row=2, min_col=len(METADATA_COLUMNS) + 1):
        for cell in row:
            cell.number_format = "@"
    # 概率单元格是逗号分隔文本，列头注释说明顺序，避免事后猜类别对应关系。
    sheet.cell(row=1, column=len(METADATA_COLUMNS) + 1).comment = Comment(
        f"每格按类别顺序 {', '.join(labels)} 给出概率，逗号分隔，合计为 1。",
        "SpecAutoAI",
    )
    headers = workbook[PREDICTION_SHEET][1]
    headers[min(len(METADATA_COLUMNS), len(headers)) - 1].comment = Comment(
        "交叉验证口径下每条记录只在它作为测试集的那一折出现（pooled OOF），因此全部记为 test；"
        "非交叉验证口径按最终模型对全部记录的预测填写 train/valid/test。",
        "SpecAutoAI",
    )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
